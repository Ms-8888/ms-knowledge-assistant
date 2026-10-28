"""Generates the DOCX and PDF seed documents (fictional content). Dev-only; needs python-docx and reportlab."""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

SEED = Path(__file__).resolve().parent.parent / "seed"

ONBOARDING = [
    ("Onboarding van nieuwe medewerkers", None),
    ("Eerste dag", "Op de eerste dag ontvang je een laptop, een badge en een rondleiding. Je teamleider stelt je voor aan het team en bespreekt je planning voor de eerste week."),
    ("Proefperiode", "De eerste maand is een proefperiode met wekelijkse feedbackgesprekken. Na vier weken volgt een evaluatiegesprek met je teamleider en HR."),
    ("Meter of peter", "Elke nieuwe medewerker krijgt een meter of peter toegewezen. Die persoon beantwoordt praktische vragen gedurende de eerste drie maanden."),
]

VEILIGHEID = [
    ("Veiligheid en preventie", None),
    ("Brand", "Bij brandalarm verlaat iedereen het gebouw via de dichtstbijzijnde nooduitgang en verzamelt op de parking aan de voorkant. Gebruik nooit de lift."),
    ("Eerste hulp", "De EHBO-koffers hangen op elke verdieping naast de keuken. De bedrijfshulpverleners zijn herkenbaar aan een groene badge. In levensgevaar bel je 112."),
    ("Bezoekers", "Bezoekers melden zich aan bij de receptie, ontvangen een bezoekersbadge en worden steeds begeleid door de medewerker die ze ontvangt."),
]

doc = Document()
for title, text in ONBOARDING:
    if text is None:
        doc.add_heading(title, level=1)
    else:
        doc.add_heading(title, level=2)
        doc.add_paragraph(text)
doc.save(SEED / "onboarding.docx")

styles = getSampleStyleSheet()
story = []
for title, text in VEILIGHEID:
    story.append(Paragraph(title, styles["Heading1" if text is None else "Heading2"]))
    if text:
        story.append(Paragraph(text, styles["BodyText"]))
    story.append(Spacer(1, 10))
SimpleDocTemplate(str(SEED / "veiligheid.pdf"), pagesize=A4, title="Veiligheid en preventie").build(story)
print("written", sorted(p.name for p in SEED.iterdir()))
